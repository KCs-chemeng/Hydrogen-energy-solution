"""Rebuild the public 100-task LNG/H2/CO2 planning workbook from sanitized schedule data."""
from pathlib import Path
from uuid import uuid5,NAMESPACE_URL
from collections import Counter
from datetime import date
import json,sys
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment

ROOT=Path(__file__).resolve().parents[1]
APP=ROOT/'Project_Schedule_Lab_v0.1.0/Project_Schedule_Lab_v0.1.0'
sys.path.insert(0,str(APP))
from model import canonical,sync_pred
from excel_io import read_excel,write_excel
from engine import deterministic,simulate,summary,textday,percentile,inspect

OUTPUT=Path(__file__).with_name('LNG_CO2_2027_01_100Tasks_Import.xlsx')
OLD=Path(__file__).with_name('LNG_CO2_2027_01_Schedule_Import.xlsx')
PACKAGE_DEFS=[
 ('P01','LNG 저장·기화 Package','LN', (18,28,45),(10,15,25),(7,10,18),(65,95,155),(10,15,25),'탱크·기화기와 연결 부속 범위. 탱크 자체를 Skid로 가정하지 않음.'),
 ('P02','가스 조압·계량·차단 Skid','NG', (12,18,32),(6,10,18),(7,10,17),(75,110,170),(8,14,26),'원료가스 공급자 분계점부터 개질기 전단까지의 공급·차단 인터페이스.'),
 ('P03','개질·수소정제 Skid','PR', (12,18,32),(6,10,18),(7,10,17),(70,105,165),(8,14,26),'개질·정제 기능을 하나의 발주 관리 Package로 가정. 최종 공급계약에서 분할 여부 확정.'),
 ('P04','수소 압축·버퍼 Package','CP/TK', (12,18,32),(6,10,18),(7,10,17),(75,110,170),(8,14,26),'압축기 Skid와 버퍼용기를 같은 인터페이스 관리 Package로 가정. 용기 검사 포함.'),
 ('P05','CO2 포집·재생 Skid','CC-capture', (14,24,44),(9,15,27),(4,7,13),(60,95,170),(8,15,28),'포집·재생·흡수제 계통. 재생열·배출·폐흡수제와 제품 경계 포함.'),
 ('P06','CO2 압축·액화 Skid','CC-liquefaction', (14,24,44),(9,15,27),(4,7,13),(60,95,170),(8,15,28),'CO2 압축·냉동·액화 계통. 저온·전력·제품 품질 인터페이스 포함.'),
 ('P07','CO2 저장·출하 Package','CC-storage', (14,24,44),(9,15,27),(4,7,13),(50,80,145),(8,15,28),'CO2 저장탱크·출하 경계. 저장탱크를 일괄 Skid로 간주하지 않음.'),
 ('P08','공통 유틸리티 Skid','UT', (12,18,32),(6,10,18),(7,10,17),(55,80,125),(8,14,26),'IA·냉각·순수·질소의 기능별 공급 범위를 집계. 외부 인입은 F03에서 관리.'),
 ('P09','전기·제어·분석 Package','PN/AN', (12,18,32),(6,10,18),(7,10,17),(60,85,135),(8,14,26),'MCC·PLC·제어반 및 분석/샘플링 Panel. 단일 업체 공급 확정이 아닌 발주관리 단위.'),
]
MODULES={
 'A':'01 착수·부지·사업조건','B':'02 공정·레이아웃·허가도서','C':'03 인허가·안전·착공조건',
 'D':'04 상세설계·시공도서','E':'14 토건·Package 설치','F':'15 현장 연결·통합 시공','G':'16 검사·시운전·인수',
}
# Tasks are deliberately rewritten as control deliverables, not formed by summing serial row durations.
BASE_ROWS=[
 ('A01','사업범위·용량·실행기준 및 일정 관리체계',3,5,8,[],'PM01 PM01A PM02 PM06','계약·운영 요구조건과 관리 기준 확정.'),
 ('A02','부지권원·접속점·현장 지장물 확인',5,10,18,['A01'],'LD01 LD04 LD03 PM03 PM04','토지·인입 통과권과 부지 제공조건 확인.'),
 ('A03','측량·현황조사 및 설계 기준점 확보',5,8,14,['A02'],'LD02A LD02B','현황·레벨·경계·지장물 조사 성과.'),
 ('A04','지반조사·기초 설계조건 확보',10,15,25,['A02'],'LD02C LD02','토질·지하수·기초형식 설계 입력 확보.'),
 ('A05','공급방식·인허가·안전심사 적용범위 확정',5,8,15,['A01','A02'],'PM05 SH01 SA04 QC01 OP01','법정 대상은 여기서 확인. PSM/기타 심사의 중복 적용을 의미하지 않음.'),
 ('B01','수소·CO2 생산량·제품규격·운전조건 정의',5,8,14,['A01'],'DB01 DB02B CB01 CB02 CB03','설계 용량과 수요·품질·배터리리밋.'),
 ('B02','LNG·원료가스 및 외부 유틸리티 공급경계 확정',5,10,18,['A02','A05','B01'],'NG00A NG00B NG00C DB02 DB02A TI01 LN01','이 초안은 현장 LNG 저장·기화를 포함. 실제 공급방식 확정 시 패키지 경계를 재검토.'),
 ('B03','PFD·물질/열·유틸리티 및 CO2 수지 확정',7,12,20,['B01','B02'],'DB03 CB04 CB05 CB10 CB11','계통·부하·CO2 회수/출하 조건과 물질수지.'),
 ('B04','예비 Vendor GA·하중·노즐·전력 자료 확보',7,12,20,['B03'],'VP01 VP00A VP00B CB07','허가용 예비 자료. 발주 후 최종 Vendor 자료와 구분.'),
 ('B05','통합 레이아웃·안전거리·동선·정비공간 검토',7,10,18,['A03','B03','B04'],'DB04 DB05 LN02','수소·LNG·CO2 구역과 건축 배치의 상호 간섭 검토.'),
 ('B06','레이아웃 동결 및 건축허가 설계입력 승인',1,3,5,['B05'],'DB06','허가용 도서 착수 기준. 계약상 승인/법정 허가를 새로 생성하는 기능 아님.'),
 ('B07','건축허가용 건축·구조·소방 도서 작성·제출본 확정',10,15,25,['B06','A04'],'BP01 CV04 CV05 CV06','허가 수준 도서를 선행 작성. 상세 IFC·최종 Vendor 앵커 도서와 분리.'),
 ('C01','권원·도시계획·실시계획 해당 행정절차',20,35,60,['A02','A05','B06'],'UP01 UP02 UP03 UP04 UP05 EX01 EX02 EX03 EX04 EX05 EX06','적용 행정절차를 하나의 관리 Package로 집계. 심사기간은 가정.'),
 ('C02','건축허가 신청·심사·보완 대응',20,30,50,['B07','C01'],'BP02 BP03','상세 설계 전체 완료를 신청 선행으로 두지 않음.'),
 ('C03','건축허가 확인·조건 배포',1,3,7,['C02'],'BP04','건축·소방 관련 허가조건과 공사 착수 조건 확인.'),
 ('C04','LNG·원료가스·수소 해당 기술검토·허가',25,40,65,['B07','D01','P01-03','P02-03','P03-03','P04-03'],'KG01 KG02 KG03 KG03A KG04 KG05 KG06 KG07 KG08 KG09 KG10','해당 시설의 확정 사양·기술검토·허가 조건 확인.'),
 ('C05','CO2 해당 기술검토·허가 및 검사 Hold 설정',25,40,65,['B07','D01','P05-03','P06-03','P07-03'],'KC01 KC02 KC03 KC04 KC05 KC06 KC07 KC08 KC09','CO2 기초·배관·완성검사 Hold와 기관 입회 조건 포함.'),
 ('C06','환경 인허가·폐기물 및 설치조건 확인',25,40,70,['A05','B03','B06','P05-02','P06-02'],'EV01 EV02 EV03 EV04 EV05 EV06 EV07 EV08 EV09 EV10 EV11 EV12','적용 분야 도서를 병행 작성한 집계 기간. 개별 기간을 직렬 합산하지 않음.'),
 ('C07','HAZOP·PSM/해당 안전심사 및 설계조치 종결',25,40,65,['A05','D01','D04','P03-03','P05-03','P06-03'],'HZ01 HZ02 HZ03 HZ04 HZ05 SH02 SH03 SH03A SH04 SH05 SH06 SH07 PS01 PS02 PS03 PS04 PS05 PS06 PS07','공정·제어 설계 조치 반영과 적용 심사조건을 집계. 대상 여부는 A05에서 확정.'),
 ('C08','시공·감리 선정 및 착공 안전계획·신고 준비',20,30,50,['B07'],'CT01 CT02 CT03 CT04 CT05 CT06 SA01 SA02 SA03 SA05 SA06 SA07','허가용 도서를 바탕으로 계약·안전·착공 신고자료를 준비.'),
 ('C09','착공 신고·해당 인허가 및 시공조건 확보',2,4,8,['C01','C03','C04','C05','C06','C07','C08','D03'],'MS01 ST01 ST02 MS02','적용 착공조건을 충족한 뒤 토건 착수. 법정 면제/승인을 가정하지 않음.'),
 ('D01','상세 P&ID·설계계산·공정안전 입력 확정',10,15,25,['B07','P03-02','P05-02','P06-02'],'PD01 PD02 PD03 PD04 CB06 CB08 CB09 CB12','건축허가용 제출본 이후 착수. Vendor 초기자료를 반영하며 심사 보완은 C07에 포함.'),
 ('D02','외부 공정·가스·CO2 배관 ISO·BOM·IFC',15,25,40,['B07','D01','P02-03','P03-03','P04-03','P05-03','P06-03','P07-03','P08-03'],'PD05 PD06 PD07 PDF1 PD08','최종 노즐·연결점·재질·압력조건을 시공 도서에 반영.'),
 ('D03','토건·장비기초·앵커·매립물 IFC',12,20,35,['B07','A04']+[f'P{i:02d}-03' for i in range(1,10)],'CV01 CV02 CV03 CV05 CVF1 CV07 CI01','허가용 도서와 분리. 최종 장비 하중·앵커·설치조건을 반영.'),
 ('D04','계장 C&E·I/O·방폭·시퀀스 상세도서',12,20,35,['B07','D01','P03-03','P05-03','P06-03','P08-03','P09-03'],'IC01 IC02 IC03 IC04 IC05 IC06 ICF1 IC07 CB08','패키지 인터페이스를 반영. HAZOP/심사 피드백 종결은 C07에서 관리.'),
 ('D05','전기·수전·통신·소방 상세도서 IFC',12,20,35,['B07','P09-03']+[f'P{i:02d}-02' for i in range(1,9)],'EL01 EL02 EL03 EL04 EL05 EL06 ELF1 EL07','최종 부하·전원·접지·방폭·소방 조건 반영.'),
 ('E01','토공·공통기초·기초검사·타설·양생',30,45,70,['C09','D03'],'CN00 CN01 CN02 CN03 CN04 CN05 CN05A CN06 CN07 CN07A CI02 CI03 CI04 BW00 UF00 UF03A UF01A UF02 UF02A','CO2 등 적용 기초 중간검사 Hold 해제 후 타설, 강도·치수 인수까지 포함.'),
 ('E02','건축·방호벽·전기실 및 설치 작업면 인계',25,40,65,['E01','D05'],'BW01 BW02 BW03 BW04 BW05 BW06 UF01 UF03 UF04 UF05 DR01 DR02 DR03 DR04 MS03 MS03E MS03U MS03P','각 구역 공사와 작업면 인계를 집계. 공통 토건 완료 후 패키지별 설치를 병행.'),
 ('F01','외부 공정·원료가스·CO2 배관 및 Skid Tie-in',21,30,45,['D02','E02'],'MP01 MP02 MP02A MP03 MP04 MP05A MP05B NG08 NG09 TI02A TI02','작업면 인계 후 선행 배관과 순차 Tie-in. SS 착수·FF 완료 조건을 관계 시트에서 확인.'),
 ('F02','현장 전력·제어·계장 케이블·접지 시공',18,28,45,['D04','D05','E02'],'EC01 EC02 EC02A EC03 EC04','Skid 설치와 병행하는 현장 배선·튜빙. 잔여 마감은 Skid 설치 완료에 연동.'),
 ('F03','수전·용수·유틸리티 인입 및 무가스 공급 준비',15,25,45,['F02','E10','E11','D05'],'PW01 PW02 PW03 EN01 EN02 TI03 NG10 NGREADY LN08 LN09','부담금·공급자 공사·전기 사용전검사와 공급 준비 집계. 가스 도입은 G07 이후.'),
 ('F04','배관 NDE·내압/기밀·세정 및 전기·루프 시험',12,20,35,['F01','F02','F03'],'MP05 MP06A MP06 MP07A MP07 MP08 MP09 EC05 EC06 EC07 CI09 CI10 CI11','현장 공정·CO2 배관 검사와 전기·제어 통합 시험을 완료.'),
 ('F05','기계적 완성·설치 검수 및 시운전 Punch 분류',5,8,15,['F04']+[f'E{i:02d}' for i in range(3,12)],'CI12 CI13 HO03A HO03B MS04','가동 차단 Punch와 잔여 항목을 구분. 무가스 SAT 준비 완료.'),
 ('G01','운영·공급계약·안전관리·운전절차 및 사전교육',15,25,40,['A05','D01','D04'],'OP02 OP03 OP04 OP05 CX01','공급·CO2 인수·운영 조직·교육 조건 확보. 담당자 개인명은 입력하지 않음.'),
 ('G02','제품 검사·인증·제조기록 및 기관 입회 준비',7,12,20,[f'P{i:02d}-05' for i in range(1,10)]+['C04','C05','C07'],'QC02 QC03 CM04A CM04B KC10','Package별 FAT와 제품/압력용기·인증 증빙 취합 및 입회 준비.'),
 ('G03','무가스 SAT·비상정지·패키지 연동 시험',7,12,20,['F05','G01'],'CM01 CM02 CM03','가스 도입 없이 제어·차단·보조계통 시험 및 결함 조치.'),
 ('G04','해당 가스시설 완성검사·보완 종결',7,12,22,['G03','G02','C04','C05'],'CM04 CM05 KC11','수소·원료가스·LNG·CO2 적용 검사/지적사항 종결.'),
 ('G05','건축·소방 준공검사·사용승인 및 가동 행정',12,20,35,['E02','F04','C03','C06'],'BU00 BU01A BU01 BU02 BU03','건축/소방 등 해당 사용승인 조건 종결. 환경 가동 전 조건은 G06.'),
 ('G06','환경·PSM 설치단계 및 가동 전 이행조건 확인',7,12,22,['F05','C06','C07','G01'],'EV13 PS08 PS09 CX02','설치 단계 확인·운영 절차·환경 신고 등 가동 전 적용 조건 종결.'),
 ('G07','가스 도입 전 PSSR·가동 허용 Gate',0,0,0,['G01','G03','G04','G05','G06'],'MS05 CO00','검사·사용승인·운영·안전 조건 충족 뒤 가스 도입을 허용.'),
 ('G08','퍼지·LNG/가스 투입·수소 생산 안정화',7,12,20,['G07'],'CO01 CO02 CO03 CX03','승인 절차에 따른 퍼지·기화·가스 투입 및 수소 계통 안정화.'),
 ('G09','CO2 포집·액화·저장·출하 연동 운전',7,12,20,['G08'],'CX04 CX05 CX06 CX07 CO04 CO04A','제품 순도·수분·출하 적합성과 H2/CO2 연계 운전 확인.'),
 ('G10','수소·CO2 통합성능·연속운전·보증 시험',10,15,25,['G09'],'CO05 CO06 CX08 PS10','생산량·포집률·품질·유틸 소비·트립/재기동·적용 심사 후속 조건 확인.'),
 ('G11','운영자 실기교육·준공도서·예비품 인계',5,8,15,['G10'],'CO07 HO00 HO02A HO02 CX09','공장 제조·현장 시험·준공 도서 및 유지보수 자료 인계.'),
 ('G12','잔여 Punch·인수조건·사업비 정산 종결',7,12,22,['G10','G11'],'HO01 HO03 HO04 HO05','운영 전 행정/사업개시 조건과 잔여 인수·정산 항목 종결.'),
 ('G13','최종 인수·사업 종결 Gate',0,0,0,['G12'],'MS06 MS07','보증시험·문서·Punch·인수조건을 충족한 최종 종결.'),
]


def build_project():
    tasks=[];edges=[];origins={};packages=[]
    project_uid=str(uuid5(NAMESPACE_URL,'psl:lng-co2:2027-01:r2-100'))
    def add(k,name,a,m,b,preds,module,old_ids='',note=''):
        tasks.append(dict(id=k,uid=str(uuid5(NAMESPACE_URL,project_uid+':'+k)),wbs=k,name=name,module=module,
                          state='계획',calendar='달력일',calendarId='default',o=a,m=m,p=b,pred=[],
                          role='',deliverable='',note=note))
        origins[k]=set(old_ids.split())
        for pred in preds:edges.append(dict(pred=pred,succ=k,type='FS',lag=0,lagUnit='달력일',lagCalendar='default',note='선행 산출물/게이트 완료 후 착수',deliverable=''))
    for k,name,a,m,b,preds,old_ids,note in BASE_ROWS:
        add(k,name,a,m,b,preds,MODULES[k[0]],old_ids,note)
    labels=['RFQ·기술평가·발주','Vendor GA·하중·노즐·부하자료 제출','제작도·인터페이스·ITP 승인','자재조달·제작·조립','FAT·Punch·출하·현장 인수']
    for index,(code,name,oldprefix,*rest) in enumerate(PACKAGE_DEFS,5):
        *durations,scope=rest
        module=f'{index:02d} {name}'
        packages.append(dict(code=code,name=name,module=module,scope=scope,stages=5))
        for stage,(label,duration) in enumerate(zip(labels,durations),1):
            preds=['B06'] if stage==1 else [f'{code}-{stage-1:02d}']
            if stage==3 and code in ('P02','P03','P04','P05','P06'):preds+=['D01']
            if stage==4 and code=='P09':preds+=['D04','D05']
            note=scope+' | Package 제작도·ITP 승인 후 제작. 현장 착공·설치·가동 조건은 C09/G07에서 별도 확인. 미결 설계조치가 제작에 영향 시 해당 Package Hold를 추가.'
            add(f'{code}-{stage:02d}',name+' | '+label,*duration,preds,module,note=note)
        num=int(code[1:])+2
        install=(7,12,20) if code not in ('P01','P03','P05','P06') else (10,15,25)
        add(f'E{num:02d}',name+' | 양중·정렬·고정·설치검수',*install,['E02',f'{code}-05'],MODULES['E'],note='해당 Package 입고와 작업면 인계 후 독립 착수. 전체 Package 납품까지 일괄 대기하지 않음.')
    # Integration work overlaps Skid installation but cannot finish ahead of the required connections.
    edges.append(dict(pred='E04',succ='F01',type='SS',lag=3,lagUnit='달력일',lagCalendar='default',note='원료가스 Skid 설치 착수 3일 후 작업면 확보 가정. 실제 작업면 인계에 맞춰 수정.',deliverable=''))
    for pred in ('E03','E04','E05','E06','E07','E08','E09'):
        edges.append(dict(pred=pred,succ='F01',type='FF',lag=5,lagUnit='달력일',lagCalendar='default',note='각 Skid 설치 완료 후 Tie-in 마감 5일 확보 가정.',deliverable=''))
    edges.append(dict(pred='E11',succ='F02',type='SS',lag=2,lagUnit='달력일',lagCalendar='default',note='반 설치 착수 2일 후 현장 배선 병행 가정.',deliverable=''))
    for pred in ('E05','E06','E07','E08','E09','E10','E11'):
        edges.append(dict(pred=pred,succ='F02',type='FF',lag=3,lagUnit='달력일',lagCalendar='default',note='각 Package 설치 완료 후 단말/결선 마감 3일 확보 가정.',deliverable=''))
    tasks.sort(key=lambda t:(t['module'],t['id']))
    p=canonical(dict(id='LNG-CC-202701-R2-100',uid=project_uid,name='LNG 수소생산·탄소포집 | 100업무·Skid 발주형 R2',
                     start='2027-01-01',asof='2027-01-01',target='',modules={t['module']:'적용' for t in tasks},
                     tasks=tasks,relations=edges,holidays=[],workdays=[],calendars={},data_version=0))
    # Keep modules in workflow display order rather than lexical task order.
    p['modules']=dict(sorted(p['modules'].items()))
    sync_pred(p)
    return p,origins,packages


def make_crosswalk(old,p,origins):
    """Conceptual scope mapping: one old activity can feed several new package/interface controls."""
    valid={t['id'] for t in old['tasks']};byold={k:[] for k in valid}
    for new,ids in origins.items():
        for k in ids & valid:byold[k].append(new)
    standard={'NG':'P02','PR':'P03','CP':'P04','TK':'P04','UT':'P08','PN':'P09','AN':'P09'}
    installation={'INNG':'E04','INPR':'E05','INCP':'E06','INTK':'E06','INUT':'E10','INPN':'E11','INAN':'E11'}
    module_fallback={
        '01':'A01','02':'A02','03':'C01','04':'B03','05':'D01','06':'D02','07':'D04','08':'D05','09':'D03',
        '10':'C02','11':'C07','12':'C04','13':'C07','14':'C08','15':'C09','16':'P03-04','17':'E01',
        '18':'E02','19':'E02','20':'E02','21':'F05','22':'F04','23':'F02','24':'F03','25':'G01',
        '26':'G04','27':'G05','28':'G09','29':'G12','30':'B02','31':'P01-04','32':'D01','33':'P05-04',
        '34':'F05','35':'C05','36':'C06','37':'C07','38':'G10'}
    mapping=[]
    for t in old['tasks']:
        k=t['id'];targets=byold[k];method='성과물·조건 중심 집계/재배치'
        prefix=next((pre for pre in standard if k.startswith(pre) and k[len(pre):] in ('01','01A','02','03A','03','04A','04','05','06A','06','06B','07')),None)
        if prefix:
            suffix=k[len(prefix):];stage=1 if suffix in ('01','01A','02') else 2 if suffix=='03A' else 3 if suffix=='03' else 4 if suffix in ('04A','04','05') else 5
            targets.append(f'{standard[prefix]}-{stage:02d}');method='발주 Package와 5개 조달 단계로 재배치'
        pre=next((pre for pre in installation if k.startswith(pre)),None)
        if pre:
            targets.append('P'+{'INNG':'02','INPR':'03','INCP':'04','INTK':'04','INUT':'08','INPN':'09','INAN':'09'}[pre]+'-05' if k.endswith('00') else installation[pre]);method='Package 인수/독립 설치로 집계'
        if k.startswith('LN') and k[2:].isdigit():
            n=int(k[2:])
            if n==4:targets+=['P01-01','P01-02']
            if n==5:targets+=['P01-05']
            targets.append('B02' if n==1 else 'B05' if n==2 else 'P01-01' if n==3 else 'P01-03' if n==4 else 'P01-04' if n==5 else 'E01' if n==6 else 'E03' if n==7 else 'F03');method='LNG 공급 Package와 현장 연결/게이트로 재배치'
        if k.startswith('CC'):
            if k in ('CC01','CC02','CC03'):targets += ['P05-01','P06-01','P07-01']
            elif k=='CC04':targets += ['P05-02','P06-02','P07-02']
            elif k=='CC05':targets += ['P05-03','P06-03','P07-03']
            elif k in ('CC06','CC07','CC08'):targets.append({'CC06':'P05-04','CC07':'P06-04','CC08':'P07-04'}[k])
            else:targets += ['P05-05','P06-05','P07-05']
            method='공통 CO2 발주를 포집/액화/저장 Package로 분할'
        if k in ('CI05','CI06','CI07','CI08'):
            targets += {'CI05':['P05-05','P06-05','P07-05'],'CI06':['E07','E08','E09'],'CI07':['F01'],'CI08':['F02']}[k]
            method='CO2 입고·설치와 현장 배관/전계장으로 분리 집계'
        if k.startswith('VP') and k not in ('VP01','VP00A','VP00B'):
            targets += [f'P{i:02d}-03' for i in range(1,10)];method='전사 VP 일괄완료 대기를 Package별 승인으로 분산'
        if not targets:
            targets.append(module_fallback[t['module'][:2]]);method='동일 분야 성과물 집계(원 행과 일대일 공정 동치는 아님)'
        for new in sorted(set(targets)):
            mapping.append(dict(old=k,old_name=t['name'],new=new,method=method))
    assert {r['old'] for r in mapping}==valid
    assert {r['new'] for r in mapping}<={t['id'] for t in p['tasks']}
    return mapping


def make_workbook():
    p,origins,packages=build_project();errors,warnings,order=inspect(p);assert not errors,errors
    assert len(p['tasks'])==100
    old=read_excel(OLD);crosswalk=make_crosswalk(old,p,origins)
    modal=deterministic(p);result=simulate(p,5000,42)
    write_excel(OUTPUT,p,result)
    w=load_workbook(OUTPUT);w.properties.creator=w.properties.lastModifiedBy='ProjectScheduleLab'
    def add_sheet(name,headers,rows):
        s=w.create_sheet(name);s.append(headers)
        for row in rows:s.append(row)
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='183653')
        for row in s.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,str):c.data_type='s'
                if isinstance(c.value,date):c.number_format='yyyy-mm-dd'
                c.alignment=Alignment(vertical='top',wrap_text=True)
        for col in s.columns:s.column_dimensions[col[0].column_letter].width=26
        return s
    out=add_sheet('M기준_예측일정',['Activity ID','모듈','업무','M 예측 시작','M 예측 완료','총 여유(달력일)','임계'],[
        [t['id'],t['module'],t['name'],date.fromisoformat(textday(modal['starts'][t['id']][0])),date.fromisoformat(textday(modal['ends'][t['id']][0])),modal['float'][t['id']],'예' if t['id'] in modal['critical'] else ''] for t in p['tasks']])
    out.column_dimensions['C'].width=70
    out=add_sheet('Package별_범위',['Package','명칭','포함/인터페이스 범위','발주','Vendor자료','제작승인','제작·조립','FAT·반입','설치'],[
        [d['code'],d['name'],d['scope']]+[f"{d['code']}-{stage:02d}" for stage in range(1,6)]+[f"E{int(d['code'][1:])+2:02d}"] for d in packages])
    out.column_dimensions['C'].width=85
    out=add_sheet('구433_신100_대응',['구 Activity ID','구 업무명','신 Activity ID','집계·재배치 기준'],[[r[k] for k in ('old','old_name','new','method')] for r in crosswalk])
    out.column_dimensions['B'].width=75;out.column_dimensions['D'].width=80
    changes=[
        ['관리 수준','433개 / 38모듈','100개 / 16모듈','세부 회신·자료접수·검측 행을 성과물·Gate로 집계. 세부 체크는 해당 100개 업무 안에서 관리.'],
        ['설계 순서','일부 상세 토건도서 이후 허가용 도서','레이아웃 B06 → 허가용 도서 B07 → 상세 D01~D05','건축허가 심사 C02와 상세설계를 병행. 상세 IFC 전체를 건축허가 신청 선행으로 두지 않음.'],
        ['구매 관리','공통 구매/VP·전 분야 제작','9개 Skid/Package × 5단계 = 45개','각 Package 발주·자료·승인·제작·FAT/입고를 별도로 추적.'],
        ['설치','통합 자료/입고 대기 가능','E03~E11 Package별 독립 설치','해당 납품과 공통 작업면 인수 후 설치. 현장 배관/전계장은 SS·FF 관계로 병행/완료조건 설정.'],
        ['관계 재작성','세부 746개 관계','집계 업무 간 성과물·Gate 관계 재작성','기존 ID/관계를 단순 제거하거나 기간을 직렬 합산하지 않음. 구신 대응은 범위 추적이며 기간/공정 일대일 동치 아님.'],
        ['가동 전 조건','개별 법정·준공·운영 조건','G07 PSSR·가동허용 Gate','완성검사·사용승인·환경/안전 이행·운영조건을 가스 도입 전에 확인.'],
        ['날짜·기간','2027-01-01 / 원 433행 가정값','2027-01-01 유지 / 집계 O/M/P 재설정','달력일 기준. 계약 납기·공휴일·고정일·실적·승인 기준 일정 미지정. 기간은 신규 발주/견적 확인 전 계획 가정.'],
        ['데이터 보존','433업무 파일과 기존 DB','별도 신규 프로젝트 LNG-CC-202701-R2-100','기존 433개 데이터/이력을 덮어쓰거나 행 삭제하지 않음. 새 파일을 신규 프로젝트로 Import.'],
        ['개인정보','공개 일정용 파일','담당자 공란 / 원본 참고자료 제외','본 변경 설명·구신 ID 대응·Package 범위는 일반 일정 정보만 포함.'],
    ]
    out=add_sheet('변경요약',['항목','기존','수정','설명'],changes);out.column_dimensions['D'].width=100
    assumptions=[
        ['공급 범위','현장 LNG 저장·기화, 수소 생산·정제·압축/버퍼, CO2 포집·액화·저장·출하 및 공통 유틸/전기제어를 포함한 계산용 초안.'],
        ['Package 경계','9개 발주관리 Package는 계획 가정. Tank/용기·전기반을 물리적 Skid 하나로 단정하지 않으며 공급계약에서 경계를 확정.'],
        ['납기','RFQ/TBE·발주, Vendor 자료, 제작승인, 제작·조립, FAT·운송/인수를 분리. 제조 단계 O/M/P는 재견적 전 집계 계획값.'],
        ['원료 공급','LNG 저장·기화를 포함하는 기존 계산 범위를 유지. 공급방식이 확정되면 P01·P02와 해당 연결 조건을 재검토.'],
        ['허가/설계','허가용 예비 하중·GA와 시공용 최종 Vendor 하중·앵커를 구분. 건축허가용 도서 이후 상세 도서 착수, 심사는 병행.'],
        ['제작/현장 분리','Package 제작도·ITP 승인 후 제작을 진행하며 현장 착공·설치·가동 조건은 C09/G07에 별도 유지. 미결 설계조치가 제작 범위에 영향 시 발주 Package별 Hold를 추가.'],
        ['달력','모든 입력 기간/시차는 달력일. 공휴일은 자동 등록하지 않음. 2027-01-01은 프로젝트 기준일이며 근무일이라고 선언하는 것이 아님.'],
        ['현장 병행','작업면·시공팀 확보 가정. F01 SS+3/FF+5, F02 SS+2/FF+3은 착수·마감 여유의 명시적 계획값이며 실적 확인 후 수정.'],
        ['자원','공종·Package를 동일 담당자라는 이유로 직렬화하지 않음. 자원/작업면 부족 시 재계획 필요.'],
        ['가동 조건','인허가 적용성은 A05에서 확정. 모든 승인·검사·법정 조건을 이 파일의 일정값만으로 대체하지 않음.'],
        ['비교 해석','433→100은 관리 상세도 변경이며 공기를 단축한다는 보장이 아님. 구신 대응은 범위 중심으로 다대다일 수 있음.'],
        ['실적·기준 일정','실제 날짜와 승인 기준 일정은 공란. 예측 결과를 기준 시작/완료로 자동 전용하지 않음.'],
        ['Import','Project Schedule Lab v0.2.2 → 엑셀 가져오기 → 새 프로젝트 → 상단 선택한 변경 적용. 100개 확인 후 시뮬레이션.'],
        ['확장 시트','변경요약/Package별 범위/구신대응/작성가정은 참고용. 앱 재Export 시 참고 시트는 보존되지 않으므로 원본 파일 별도 보관.'],
    ]
    out=add_sheet('작성가정',['항목','내용'],assumptions);out.column_dimensions['B'].width=120
    stats=summary(result,p);old_stats=summary(simulate(old,5000,42),old)
    summary_rows=[['구분','기존 433개','수정 100개'],['프로젝트 시작',old['start'],p['start']],['M 기준 완료',textday(deterministic(old)['total'][0]),textday(modal['total'][0])]]
    summary_rows += [[q,old_stats[q],stats[q]] for q in ('P50','P80','P90')]
    add_sheet('예측비교',summary_rows[0],summary_rows[1:])
    w.save(OUTPUT);w.close()
    report=dict(project_id=p['id'],tasks=len(p['tasks']),modules=len(p['modules']),packages=len(packages),procurement_tasks=45,
                relations=len(p['relations']),relation_types=dict(Counter(r['type'] for r in p['relations'])),
                mapped_old_tasks=len({r['old'] for r in crosswalk}),old_tasks=len(old['tasks']),M_finish=textday(modal['total'][0]),
                forecast=stats,old_forecast=old_stats,critical=modal['critical'])
    Path(__file__).with_name('LNG_CO2_2027_100Tasks_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return p,report

if __name__=='__main__':
    p,report=make_workbook()
    print(json.dumps(report,ensure_ascii=False,indent=2))
