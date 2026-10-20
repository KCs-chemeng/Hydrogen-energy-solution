"""Acceptance checks for the revised planning data, independent of the application's source tests."""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import numpy as np
from openpyxl import load_workbook
from build_100_task_schedule import build_project,make_crosswalk,OLD,OUTPUT,read_excel
from engine import inspect,deterministic,simulate,Calendar,day
from importing import prepare
from storage import Store


class RevisedScheduleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.p,cls.origins,cls.packages=build_project()
        cls.r=simulate(cls.p,5000,42)
    def test_scope_ids_and_privacy(self):
        p=self.p
        self.assertEqual(len(p['tasks']),100);self.assertEqual(len(p['modules']),16)
        self.assertEqual(p['start'],p['asof']);self.assertEqual(p['start'],'2027-01-01')
        self.assertNotEqual(p['uid'],read_excel(OLD)['uid'])
        self.assertEqual(len({t['id'] for t in p['tasks']}),100)
        self.assertEqual(len({t['uid'] for t in p['tasks']}),100)
        self.assertEqual(inspect(p)[0],[]);self.assertEqual(set(p['modules'].values()),{'적용'})
        for t in p['tasks']:
            self.assertEqual(t['role'],'')
            for k in ('actualStart','actualEnd','baselineStart','baselineEnd','fixedStart','fixedEnd'):self.assertEqual(t[k],'')
            self.assertLessEqual(0,t['o']);self.assertLessEqual(t['o'],t['m']);self.assertLessEqual(t['m'],t['p'])
    def ancestors(self,node):
        found=set();pending=[node]
        while pending:
            k=pending.pop()
            for e in self.p['relations']:
                if e['succ']==k and e['pred'] not in found:found.add(e['pred']);pending.append(e['pred'])
        return found
    def test_layout_permit_drawings_then_detail_without_waiting_for_full_design(self):
        self.assertIn('B06',self.ancestors('B07'))
        for k in ('D01','D02','D03','D04','D05'):self.assertIn('B07',self.ancestors(k))
        for k in ('D01','D02','D03','D04','D05'):self.assertNotIn(k,self.ancestors('C02'))
        self.assertNotIn('C03',self.ancestors('D01'))
        for k in ('D01','D02','D03','D04','D05'):self.assertTrue(np.all(self.r['starts'][k]>=self.r['ends']['B07']))
    def test_nine_packages_five_steps_and_independent_installation(self):
        p=self.p
        self.assertEqual(len(self.packages),9)
        self.assertEqual(sum(t['id'].startswith('P') for t in p['tasks']),45)
        for i in range(1,10):
            code=f'P{i:02d}'
            for stage in range(2,6):self.assertIn(f'{code}-{stage-1:02d}',self.ancestors(f'{code}-{stage:02d}'))
            install=f'E{i+2:02d}'
            direct={e['pred'] for e in p['relations'] if e['succ']==install}
            self.assertEqual(direct,{'E02',f'{code}-05'})
        a=deterministic(p);changed=deepcopy(p)
        t=next(t for t in changed['tasks'] if t['id']=='P01-04')
        for k in ('o','m','p'):t[k]+=180
        b=deterministic(changed)
        self.assertGreater(int(b['starts']['E03'][0]),int(a['starts']['E03'][0]))
        for k in ('P02-01','P03-04','P05-04','E04'):
            np.testing.assert_equal(a['starts'][k],b['starts'][k])
    def test_all_relations_hold_in_every_sample(self):
        p,r=self.p,self.r
        self.assertTrue(any(e['type']=='SS' for e in p['relations']))
        self.assertTrue(any(e['type']=='FF' for e in p['relations']))
        for e in p['relations']:
            source=r['starts' if e['type']=='SS' else 'ends'][e['pred']]
            target=r['ends' if e['type']=='FF' else 'starts'][e['succ']]
            bound=Calendar(p,e['lagUnit'],e['lagCalendar']).add(source,int(e['lag']))
            self.assertTrue(np.all(target>=bound),str(e))
    def test_gas_introduction_and_terminal_gate(self):
        for k in ('G03','G04','G05','G06','G01'):self.assertIn(k,self.ancestors('G07'))
        self.assertIn('G07',self.ancestors('G08'))
        self.assertEqual(self.ancestors('G13'),{t['id'] for t in self.p['tasks']} - {'G13'})
        np.testing.assert_equal(self.r['total'],self.r['ends']['G13'])
    def test_constant_omp_shared_engine_and_repeatability(self):
        p=deepcopy(self.p)
        for t in p['tasks']:t['o']=t['p']=t['m']
        a=deterministic(p);b=simulate(p,64,42)
        for k in a['starts']:
            self.assertTrue(np.all(a['starts'][k][0]==b['starts'][k]));self.assertTrue(np.all(a['ends'][k][0]==b['ends'][k]))
        np.testing.assert_equal(self.r['total'],simulate(self.p,5000,42)['total'])
    def test_old_scope_is_mapped_with_no_original_deletion(self):
        old=read_excel(OLD);mapping=make_crosswalk(old,self.p,self.origins)
        self.assertEqual({r['old'] for r in mapping},{t['id'] for t in old['tasks']})
        self.assertEqual(len({r['old'] for r in mapping}),433)
        self.assertLessEqual({r['new'] for r in mapping},{t['id'] for t in self.p['tasks']})
    def test_real_workbook_import_noop_and_old_project_preserved(self):
        with TemporaryDirectory() as td:
            s=Store(Path(td)/'test.sqlite3');old=read_excel(OLD);s.add(old,'신규','기존 433개 보존 시험')
            plan=prepare(OUTPUT,None,s);self.assertFalse(plan.issues,plan.report())
            p=s.apply_import(plan);self.assertEqual(len(p['tasks']),100);self.assertEqual(len(s.projects()),2)
            self.assertEqual(len(s.load(old['uid'])['tasks']),433)
            again=prepare(OUTPUT,p,s);self.assertEqual(again.changes,[])
            version=p['data_version'];p=s.apply_import(again,acknowledge=True);self.assertEqual(p['data_version'],version)
            s.db.close()
    def test_workbook_schedule_only_and_reference_traceability(self):
        w=load_workbook(OUTPUT)
        self.assertEqual(w['작업목록'].max_row,101);self.assertEqual(w['모듈'].max_row,17)
        self.assertEqual(w['Package별_범위'].max_row,10)
        self.assertTrue({'변경요약','작성가정','구433_신100_대응','예측비교'}<=set(w.sheetnames))
        self.assertFalse({'인천업무 근거','확인사항','병행검토'} & set(w.sheetnames))
        for sh in w:
            for row in sh:
                for c in row:self.assertIsNone(c.comment);self.assertIsNone(c.hyperlink)
        w.close()

if __name__=='__main__':unittest.main(verbosity=2)
