"""Standard PSL-2 workbook, with cell-level diagnostics and legacy Rev1 input support."""
from dataclasses import dataclass, field, asdict
from datetime import date, datetime
from pathlib import Path
import json
import os
import tempfile
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from engine import day, textday, percentile, summary
from model import canonical, TASK_DEFAULTS, uid
from legacy_excel import read_excel as read_legacy, COLS as LEGACY_COLS

COLS = dict(LEGACY_COLS, wbs='WBS', calendarId='달력 ID', fixedStart='고정 시작', fixedEnd='고정 완료',
            resume='재개 가능일', progress='진척률', uid='내부 UID')
DATES = ['actualStart','actualEnd','constraint','baselineStart','baselineEnd','fixedStart','fixedEnd','resume']
NUMBERS = ['o','m','p','ro','rm','rp','progress']
CLEAR = '__CLEAR__'
REL_COLS = dict(pred='선행 ID',succ='후행 ID',type='관계 유형',lag='시차',lagUnit='시차 단위',
                lagCalendar='시차 달력',note='설정 근거',deliverable='필요 산출물')


@dataclass
class Issue:
    severity: str
    file: str
    sheet: str
    row: int
    column: str
    value: object
    message: str
    fix: str

    def __str__(self):
        return f'{self.severity} {Path(self.file).name} / {self.sheet}!{self.column}{self.row}: {self.message} (값={self.value!r}) → {self.fix}'


@dataclass
class Document:
    project: dict
    filename: str
    issues: list = field(default_factory=list)
    cells: dict = field(default_factory=dict)
    export_version: object = None
    schema: str = 'legacy'
    explicit_relations: bool = False
    metadata: dict = field(default_factory=dict)


def _date(v):
    if isinstance(v,(date,datetime)): v=v.strftime('%Y-%m-%d')
    if v in (None,''): return ''
    day(str(v).strip())
    return str(v).strip()


def read_document(path):
    path=Path(path)
    w=load_workbook(path,data_only=False)
    issues=[]; cells={}; clear_sections=[]
    def issue(s,r,c,v,message,fix,severity='오류'):
        issues.append(Issue(severity,str(path),s,r,c,v,message,fix))
    def convert(v,k,s,row,col):
        if v == CLEAR: return CLEAR
        if v in (None,''): return None if k in NUMBERS else ([] if k=='pred' else '')
        try:
            if k in NUMBERS:
                if isinstance(v,bool): raise ValueError()
                result=float(v)
                import math
                if not math.isfinite(result): raise ValueError()
                return result
            if k in DATES: return _date(v)
            if k=='pred': return [s.strip() for s in str(v).replace(';',',').split(',') if s.strip()]
            return str(v).strip()
        except (TypeError,ValueError):
            issue(s,row,col,v,'자료형/날짜 오류 또는 수식','수식 대신 유한 숫자 또는 YYYY-MM-DD 입력')
            return None if k in NUMBERS else ''
    try:
        if '작업목록' not in w.sheetnames:
            if not all(s in w.sheetnames for s in ('상세일정','모듈선택','프로젝트설정')):
                raise ValueError('지원 형식: PSL 교환 XLSX / 공통 Rev1')
            # Validate raw legacy dates/numbers before conversion; never silently repair them.
            for row in w['상세일정'].iter_rows(min_row=6):
                if not any(c.value is not None for c in row[:29]): continue
                if row[0].value in (None,''):
                    issue('상세일정',row[0].row,'A','', 'Activity ID 없음','신규 등록 시 ID를 지정하세요')
                for i,k in [(7,'baselineStart'),(8,'baselineEnd'),(9,'actualStart'),(10,'actualEnd'),(11,'constraint')]+list(zip(range(12,18),NUMBERS[:6])):
                    if len(row)>i: convert(row[i].value,k,'상세일정',row[i].row,row[i].column_letter)
            try: p=canonical(read_legacy(path),identities=False)
            except Exception as exc:
                issue('상세일정',1,'A','',str(exc),'원본을 확인하고 교환 형식으로 입력하세요')
                p=canonical(dict(id='PJ-NEW',name=path.stem,start='',asof='',tasks=[],modules={}),identities=False)
            # Locations for semantic diagnostics.
            for row in w['상세일정'].iter_rows(min_row=6):
                if row[0].value is not None:
                    k=str(row[0].value)
                    fields={0:'id',1:'module',2:'name',4:'state',7:'baselineStart',8:'baselineEnd',9:'actualStart',10:'actualEnd',11:'constraint',12:'o',13:'m',14:'p',15:'ro',16:'rm',17:'rp',23:'pred'}
                    for col,key in fields.items():
                        cells[(k,key)]=('상세일정',row[0].row,row[col].column_letter,row[col].value)
            for row in w['모듈선택'].iter_rows(min_row=6):
                if row[0].value: cells[('$module',str(row[0].value))]=('모듈선택',row[0].row,'B',row[1].value)
            for field,row in [('start',7),('asof',8),('target',11)]:
                cells[('$project',field)]=('프로젝트설정',row,'B',w['프로젝트설정'].cell(row,2).value)
            return Document(p,str(path),issues,cells,schema='Rev1')
        required_sheets=('프로젝트','모듈','작업목록')
        for s in required_sheets:
            if s not in w.sheetnames: issue(s,1,'A','', '필수 시트 없음','원본 표준 템플릿에서 복원')
        if issues: return Document({},str(path),issues)
        cfg={str(r[0].value):r[1].value for r in w['프로젝트'].iter_rows(min_row=2) if r[0].value and len(r)>1}
        schema=cfg.get('형식버전','PSL-0.1')
        if schema not in ('PSL-0.1','PSL-2'): issue('프로젝트',2,'B',schema,'미지원 스키마 버전','PSL-2 또는 PSL-0.1 사용')
        p=dict(id=str(cfg.get('PJ_ID') or ''),name=str(cfg.get('프로젝트명') or ''),
               start='',asof='',target='',modules={},tasks=[],holidays=[],workdays=[],calendars={})
        for k,label in [('start','시작일'),('asof','기준일'),('target','목표일')]:
            value=cfg.get(label)
            row=next((r[0].row for r in w['프로젝트'].iter_rows(min_row=2) if r[0].value==label),1)
            p[k]=convert(value,'constraint','프로젝트',row,'B')
            cells[('$project',k)]=('프로젝트',row,'B',value)
        if cfg.get('프로젝트 UID'): p['uid']=str(cfg['프로젝트 UID'])
        try: p['data_version']=int(cfg.get('데이터 버전') or 0)
        except (ValueError,TypeError): issue('프로젝트',1,'B',cfg.get('데이터 버전'),'잘못된 데이터 버전','Export 메타데이터를 수정하지 마세요')
        for row in w['모듈'].iter_rows(min_row=2):
            if row[0].value:
                name=str(row[0].value)
                if name in p['modules']: issue('모듈',row[0].row,'A',name,'모듈 중복','중복 행 제거')
                p['modules'][name]=str(row[1].value or '검토 필요')
                cells[('$module',name)]=('모듈',row[0].row,'B',row[1].value)
        for sheet,key in [('휴일','holidays'),('추가근무일','workdays')]:
            if sheet in w:
                for row in w[sheet].iter_rows(min_row=2):
                    if row[0].value == CLEAR:
                        clear_sections.append(key)
                    elif row[0].value: p[key].append(convert(row[0].value,'constraint',sheet,row[0].row,'A'))
        s=w['작업목록']; headers=[c.value for c in s[1]]
        for label in LEGACY_COLS.values():
            if label not in headers: issue(s.title,1,'A',label,'필수 열 누락','기존 필수 열 이름 복원')
        if len([h for h in headers if h])!=len(set(h for h in headers if h)): issue(s.title,1,'A',headers,'중복 열 이름','열 이름을 고유하게 지정')
        for row in s.iter_rows(min_row=2):
            if not any(c.value is not None for c in row): continue
            t={}
            for k,label in COLS.items():
                if label not in headers: continue
                c=row[headers.index(label)]
                t[k]=convert(c.value,k,s.title,c.row,c.column_letter)
                cells[(str(row[headers.index('Activity ID')].value or ''),k)]=(s.title,c.row,c.column_letter,c.value)
            p['tasks'].append(t)
        explicit='관계' in w
        if explicit:
            p['relations']=[]
            s=w['관계']; headers=[c.value for c in s[1]]
            for label in REL_COLS.values():
                if label not in headers: issue(s.title,1,'A',label,'관계 열 누락','관계 시트 헤더 복원')
            for row in s.iter_rows(min_row=2):
                if not any(c.value is not None for c in row): continue
                r={}
                for k,label in REL_COLS.items():
                    if label not in headers: continue
                    c=row[headers.index(label)];v=c.value
                    if k=='lag':
                        try: v=float(v or 0)
                        except (ValueError,TypeError): issue(s.title,c.row,c.column_letter,v,'시차 숫자 오류','0 이상의 정수 입력');v=-1
                    else: v=str(v or '').strip()
                    r[k]=v
                    cells[('$relation'+str(len(p['relations'])),k)]=(s.title,c.row,c.column_letter,c.value)
                p['relations'].append(r)
            # Explicit relation sheet is authoritative; conflict always reported.
            for t in p['tasks']:
                expected=set(r['pred'] for r in p['relations'] if r['succ']==t.get('id'))
                if set(t.get('pred',[]))!=expected:
                    loc=cells.get((t.get('id'),'pred'),('작업목록',1,'K',''))
                    issue(*loc,'관계 시트와 선행 ID 열 불일치: 관계 시트 우선','관계 시트를 검토 후 경고 확인',severity='경고')
        if '달력' in w:
            for row in w['달력'].iter_rows(min_row=2):
                if not row[0].value: continue
                key,typ,v=[c.value for c in row[:3]]
                cfgcal=p['calendars'].setdefault(str(key),dict(weekmask=[0,1,2,3,4],holidays=[],workdays=[]))
                if typ=='weekmask':
                    try: cfgcal['weekmask']=[int(x) for x in str(v).split(',')]
                    except ValueError: issue('달력',row[0].row,'C',v,'요일 목록 오류','월=0~일=6, 쉼표로 구분')
                elif typ in ('holidays','workdays'): cfgcal[typ].append(convert(v,'constraint','달력',row[0].row,'C'))
                else: issue('달력',row[0].row,'B',typ,'달력 항목 오류','weekmask / holidays / workdays 사용')
        metadata={'clear_sections':clear_sections}
        if '메타데이터' in w:
            metadata.update({str(r[0]):r[1] for r in w['메타데이터'].iter_rows(min_row=2,values_only=True) if r[0]})
        return Document(p,str(path),issues,cells,cfg.get('데이터 버전'),schema,explicit,metadata)
    finally: w.close()


def read_excel(path):
    doc=read_document(path)
    if any(i.severity=='오류' for i in doc.issues): raise ValueError('\n'.join(map(str,doc.issues)))
    return canonical(doc.project)


def write_excel(path,p,r=None,baselines=None):
    p=canonical(p)
    w=Workbook();w.remove(w.active)
    def sheet(name,headers,rows):
        s=w.create_sheet(name);s.append(headers)
        for row in rows: s.append(row)
        return s
    sheet('프로젝트',['항목','값'],[(k,v) for k,v in [
        ('형식버전','PSL-2'),('PJ_ID',p['id']),('프로젝트 UID',p['uid']),('데이터 버전',p['data_version']),
        ('프로젝트명',p['name']),('시작일',p['start']),('기준일',p['asof']),('목표일',p.get('target',''))]])
    sheet('모듈',['모듈','적용여부'],p['modules'].items())
    sheet('작업목록',list(COLS.values()),[[', '.join(t.get(k,[])) if k=='pred' else t.get(k) for k in COLS] for t in p['tasks']])
    sheet('관계',list(REL_COLS.values()),[[r.get(k) for k in REL_COLS] for r in p['relations']])
    sheet('휴일',['휴일 날짜'],[[v] for v in p['holidays']])
    sheet('추가근무일',['추가 근무 날짜'],[[v] for v in p['workdays']])
    rows=[]
    for k,cfg in p['calendars'].items():
        rows.append([k,'weekmask',','.join(map(str,cfg.get('weekmask',[0,1,2,3,4])))])
        for typ in ('holidays','workdays'):
            rows.extend([k,typ,v] for v in cfg.get(typ,[]))
    sheet('달력',['달력 ID','종류','값'],rows)
    sheet('메타데이터',['항목','값'],[('schema','PSL-2'),('export_version',p['data_version']),('project_uid',p['uid'])])
    if baselines:
        rows=[]
        for b in baselines:
            if b.get('legacy'): continue
            for k,s in b['payload']['starts'].items(): rows.append([b['id'],b['name'],b['version'],b['created'],k,textday(s),textday(b['payload']['ends'][k])])
        sheet('기준일정버전',['버전 ID','이름','데이터 버전','승인 시각','Activity ID','시작','완료'],rows)
    sheet('사용안내',['항목','설명'],[
        ('빈 셀','기존 데이터 보존. 명시적 삭제는 __CLEAR__. 신규 작업의 필수값 삭제는 오류.'),
        ('관계','관계 시트가 선행 ID 열보다 우선. 불일치 시 경고. 관계 변경은 미리보기에서 선택하여 적용.'),
        ('ID','Activity ID는 프로젝트 내 고유. 내부 UID와 Export 데이터 버전은 수정하지 마세요.'),
        ('달력','기본 월~금. 추가 근무일이 휴일보다 우선. 공휴일 자동 등록 없음. 1900~2199년 지원.'),
        ('일수','경과일: 10/1 시작 1일=10/2 완료. FS+0은 선행 완료일에 착수 가능. 0일은 이동 없음.'),
        ('시차','관계별 단위/달력 사용. 시차는 0 이상 정수. 근무일 양수 시차는 다음 근무일로 보정 후 더함.'),
        ('실적','기준일은 당일 시작 시점. 진행 잔여기간은 기준일부터. 미완료 FS는 잔여 작업 재개를 지연.'),
        ('내보내기','표준 데이터만 출력. 원본 임의 시트/그림/서식은 복제하지 않음. 기준 일정 시트는 조회용.'),
        ('여유','총 여유는 달력일 단위. 실적이 고정된 작업은 정의하지 않음. 목표일 차이와 별개.')])
    if r is not None:
        sheet('분석요약',['항목','결과'],list(summary(r,p).items()))
        rows=[]
        for t in p['tasks']:
            k=t['id']
            rows.append([k,t['name'],textday(percentile(r['starts'][k],50)),textday(percentile(r['ends'][k],50)),
                         textday(percentile(r['ends'][k],80)),r.get('float',{}).get(k),'예' if k in r.get('critical',[]) else ''])
        sheet('작업결과',['Activity ID','작업명','시작 P50','완료 P50','완료 P80','총 여유(달력일)','임계'],rows)
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='183653')
        for row in s.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,str): c.data_type='s'
        for col in s.columns: s.column_dimensions[col[0].column_letter].width=23
        if s.title=='작업목록': s.column_dimensions['C'].width=45
        if s.title=='사용안내': s.column_dimensions['B'].width=110
    dest=Path(path)
    fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=dest.parent);os.close(fd)
    try: w.save(tmp);os.replace(tmp,dest)
    finally:
        w.close()
        if os.path.exists(tmp): os.unlink(tmp)
