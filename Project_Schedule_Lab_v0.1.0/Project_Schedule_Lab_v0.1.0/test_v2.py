"""Data safety and schedule examples, with hand-calculated reference dates."""
import json
import sqlite3
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
import numpy as np
from openpyxl import load_workbook
from model import canonical, clone_project, content
from engine import day,textday,simulate,deterministic,inspect,Calendar
from storage import Store
from excel_io import write_excel,read_excel,read_document,COLS,CLEAR
from importing import prepare,materialize


def project(durations=(2,3), edges=None, unit='달력일'):
    tasks=[dict(id=chr(65+i),module='M',name='작업 '+chr(65+i),o=d,m=d,p=d,pred=[],state='계획',calendar=unit) for i,d in enumerate(durations)]
    p=canonical(dict(id='PJ',name='테스트',start='2026-10-01',asof='2026-10-01',target='',modules={'M':'적용'},tasks=tasks))
    p['relations']=edges or []
    from model import sync_pred
    sync_pred(p)
    return p


def edge(a,b,typ='FS',lag=0,unit='달력일',cal='default'):
    return dict(pred=a,succ=b,type=typ,lag=lag,lagUnit=unit,lagCalendar=cal,note='시험 근거',deliverable='문서')


class ScheduleTests(unittest.TestCase):
    def test_fs_and_parallel_merge(self):
        p=project((2,5,1),[edge('A','C'),edge('B','C')]);r=deterministic(p)
        self.assertEqual(int(r['ends']['C'][0]),day('2026-10-07'))
        self.assertEqual(r['float'],{'C':0,'B':0,'A':3})
        self.assertEqual(set(r['critical']),{'B','C'})
        self.assertEqual(int(r['starts']['A'][0]),int(r['starts']['B'][0]))

    def test_ss_ff_multiple_lag(self):
        p=project((4,2,1),[edge('A','B','SS',1),edge('A','C','FF',2),edge('B','C')])
        r=deterministic(p)
        self.assertEqual(int(r['starts']['B'][0]),day('2026-10-02'))
        self.assertEqual(int(r['starts']['C'][0]),day('2026-10-06'))
        self.assertEqual(int(r['ends']['C'][0]),day('2026-10-07'))
        self.assertEqual(r['float']['B'],2)
        self.assertEqual(set(r['critical']),{'A','C'})

    def test_lag_workdays(self):
        p=project((1,1),[edge('A','B','FS',1,'근무일')]);p['holidays']=['2026-10-05']
        r=deterministic(p)
        self.assertEqual(int(r['starts']['B'][0]),day('2026-10-06'))

    def test_holidays_extra_workdays_mixed_calendar(self):
        p=project((1,1),[edge('A','B')],unit='근무일');p['start']=p['asof']='2026-10-02'
        p['holidays']=['2026-10-05'];p['workdays']=['2026-10-03']
        p['calendars']={'seven':dict(weekmask=list(range(7)),holidays=[],workdays=[])}
        p['tasks'][1]['calendarId']='seven'
        r=deterministic(p)
        self.assertEqual(int(r['ends']['A'][0]),day('2026-10-03'))
        self.assertEqual(int(r['ends']['B'][0]),day('2026-10-04'))

    def test_workday_ff_and_float(self):
        p=project((4,1,1),[edge('A','C','FF'),edge('B','C')],unit='근무일')
        r=deterministic(p)
        self.assertEqual(int(r['ends']['C'][0]),day('2026-10-07'))
        self.assertEqual(r['float']['A'],0)
        self.assertEqual(r['float']['B'],4)  # Thu -> Mon: elapsed days, not workdays.

    def test_one_day_and_zero_milestone(self):
        p=project((0,1),[edge('A','B')],unit='근무일');p['start']=p['asof']='2026-10-03'
        r=deterministic(p)
        self.assertEqual(int(r['ends']['A'][0]),day('2026-10-03'))
        self.assertEqual(int(r['starts']['B'][0]),day('2026-10-05'))
        self.assertEqual(int(r['ends']['B'][0]),day('2026-10-06'))

    def test_completed_actual_violations_warn(self):
        p=project((5,1),[edge('A','B')]);p['asof']='2026-10-10'
        for t,s,f in [(p['tasks'][0],'2026-10-01','2026-10-06'),(p['tasks'][1],'2026-10-02','2026-10-03')]: t.update(state='완료',actualStart=s,actualEnd=f)
        r=deterministic(p)
        self.assertEqual(int(r['starts']['B'][0]),day('2026-10-02'))
        self.assertTrue(r['warnings']);self.assertIsNone(r['float']['B'])

    def test_ongoing_remaining_waits_for_incomplete_fs(self):
        p=project((5,10),[edge('A','B')]);p['asof']='2026-10-03'
        p['tasks'][1].update(state='진행',actualStart='2026-10-02',ro=2,rm=2,rp=2)
        r=deterministic(p)
        self.assertEqual(int(r['starts']['B'][0]),day('2026-10-02'))
        self.assertEqual(int(r['ends']['B'][0]),day('2026-10-10'))
        self.assertTrue(r['warnings'])

    def test_hold_and_missing_remaining(self):
        p=project((2,));p['tasks'][0]['state']='보류'
        with self.assertRaisesRegex(ValueError,'재개'): deterministic(p)
        p['tasks'][0]['resume']='2026-10-10';self.assertEqual(int(deterministic(p)['ends']['A'][0]),day('2026-10-12'))
        p['tasks'][0].update(state='진행',actualStart='2026-10-01',progress=50)
        with self.assertRaisesRegex(ValueError,'잔여'): deterministic(p)

    def test_fixed_dates_and_target(self):
        p=project((2,3),[edge('A','B')]);p['target']='2026-10-04'
        r=deterministic(p);self.assertEqual(r['target_delta'],-2)
        p['tasks'][1]['fixedStart']='2026-10-02'
        with self.assertRaisesRegex(ValueError,'고정'): deterministic(p)
        p['tasks'][1]['fixedStart']='2026-10-04';p['tasks'][1]['fixedEnd']='2026-10-06'
        with self.assertRaisesRegex(ValueError,'충돌'): deterministic(p)
        p['tasks'][1]['fixedEnd']='2026-10-07'
        self.assertEqual(int(deterministic(p)['ends']['B'][0]),day('2026-10-07'))

    def test_cycles_self_cross_project_and_negative_lag(self):
        p=project(edges=[edge('A','B'),edge('B','A')])
        self.assertTrue(any('A → B → A' in e for e in inspect(p)[0]))
        p['relations']=[edge('A','A')];self.assertTrue(any('자기' in e for e in inspect(p)[0]))
        p['relations']=[edge('OTHER/A','B')];self.assertTrue(any('다른 프로젝트' in e for e in inspect(p)[0]))
        p['relations']=[edge('A','B',lag=-1)];self.assertTrue(any('시차' in e for e in inspect(p)[0]))
        p['relations']=[dict(edge('A','B'),project_uid='another')];self.assertTrue(any('다른 프로젝트' in e for e in inspect(p)[0]))

    def test_exclusion_restore_preserves_network(self):
        p=project(edges=[edge('A','B')]);before=deepcopy(p['relations']);p['modules']['M']='제외'
        with self.assertRaisesRegex(ValueError,'제외'): deterministic(p)
        self.assertEqual(p['relations'],before)
        p['modules']['M']='적용';self.assertEqual(int(deterministic(p)['total'][0]),day('2026-10-06'))

    def test_fixed_omp_monte_carlo_equals_deterministic(self):
        p=project((3,2,0),[edge('A','B','SS',1),edge('A','C','FF',2)],unit='근무일')
        p['holidays']=['2026-10-05'];p['workdays']=['2026-10-03']
        a=deterministic(p);b=simulate(p,200,123)
        for k in a['starts']:
            self.assertTrue(np.all(a['starts'][k][0]==b['starts'][k]));self.assertTrue(np.all(a['ends'][k][0]==b['ends'][k]))
        self.assertEqual(b['data_version'],p['data_version']);self.assertEqual(b['asof'],p['asof'])

    def test_duration_validation_and_state_contradiction(self):
        p=project((1,));p['tasks'][0]['o']=-1
        with self.assertRaisesRegex(ValueError,'O'): simulate(p,10)
        p['tasks'][0].update(o=1,actualStart='2026-10-01')
        with self.assertRaisesRegex(ValueError,'모순'): deterministic(p)
        p['tasks'][0].update(actualStart='',progress=50)
        with self.assertRaisesRegex(ValueError,'모순'): deterministic(p)

    def test_generalized_cpm_against_enumerated_feasible_schedules(self):
        # Independent, scalar calendar arithmetic; enumerate all schedules that meet the earliest finish.
        import itertools,random
        rng=random.Random(19)
        for case in range(12):
            ds=tuple(rng.randrange(3) for _ in range(3))
            es=[edge('A','B',rng.choice(['FS','SS','FF']),rng.randrange(2),rng.choice(['달력일','근무일'])),
                edge('B','C',rng.choice(['FS','SS','FF']),rng.randrange(2)),
                edge('A','C',rng.choice(['SS','FF']),0)]
            p=project(ds,es)
            for t in p['tasks']: t['calendar']=rng.choice(['달력일','근무일'])
            p['holidays']=['2026-10-05'];p['workdays']=['2026-10-03']
            base=day(p['start'])
            def working(d): return d==day('2026-10-03') or ((d-1)%7<5 and d!=day('2026-10-05'))
            def advance(d,n,unit):
                if unit=='달력일' or n==0: return d+n
                while not working(d): d+=1
                for _ in range(n):
                    d+=1
                    while not working(d): d+=1
                return d
            r=deterministic(p);finish=int(r['total'][0]);keys=['A','B','C']
            maxstart={k:base for k in keys};found=False
            candidates=[]
            for i,t in enumerate(p['tasks']):
                candidates.append([s for s in range(base,finish+1) if (not ds[i] or t['calendar']=='달력일' or working(s)) and advance(s,ds[i],t['calendar'])<=finish])
            for ss in itertools.product(*candidates):
                starts=dict(zip(keys,ss));ends={k:advance(starts[k],ds[i],p['tasks'][i]['calendar']) for i,k in enumerate(keys)}
                feasible=True
                for e in es:
                    a,b=e['pred'],e['succ'];anchor=starts[a] if e['type']=='SS' else ends[a]
                    dependent=ends[b] if e['type']=='FF' else starts[b]
                    if dependent<advance(anchor,e['lag'],e['lagUnit']): feasible=False;break
                if feasible:
                    found=True
                    for k in keys: maxstart[k]=max(maxstart[k],starts[k])
            self.assertTrue(found)
            self.assertEqual(r['float'],{k:maxstart[k]-int(r['starts'][k][0]) for k in keys},f'case {case} {es} {ds}')


class DataTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.store=Store(self.root/'db.sqlite3');self.p=project(edges=[edge('A','B')]);self.store.add(self.p,'생성','테스트')
        self.file=self.root/'test.xlsx';write_excel(self.file,self.p)
    def tearDown(self):
        self.store.db.close();self.tmp.cleanup()
    def edit(self,fn):
        w=load_workbook(self.file);fn(w);w.save(self.file);w.close()
    def cell(self,w,field,row=2):
        return w['작업목록'].cell(row,list(COLS).index(field)+1)
    def apply(self,ack=True,selected=None):
        plan=prepare(self.file,self.p,self.store)
        p=self.store.apply_import(plan,selected,acknowledge=ack)
        self.p=p;return plan,p

    def test_noop_roundtrip_no_history_version(self):
        before=self.store.db.execute('SELECT count(*) FROM changes').fetchone()[0];events=len(self.store.rows());v=self.p['data_version']
        plan,p=self.apply()
        self.assertEqual(plan.changes,[]);self.assertEqual(p['data_version'],v)
        self.assertEqual(len(self.store.rows()),events)
        self.assertEqual(self.store.db.execute('SELECT count(*) FROM changes').fetchone()[0],before)
        self.assertEqual(content(read_excel(self.file)),content(p))

    def test_name_wbs_and_row_order_keep_ids(self):
        def edit(w):
            self.cell(w,'name').value='이름 변경';self.cell(w,'wbs').value='9.2'
            s=w['작업목록'];a=[c.value for c in s[2]];b=[c.value for c in s[3]]
            for j,v in enumerate(b,1): s.cell(2,j,v)
            for j,v in enumerate(a,1): s.cell(3,j,v)
        self.edit(edit);uids={t['id']:t['uid'] for t in self.p['tasks']};_,p=self.apply()
        self.assertEqual(uids,{t['id']:t['uid'] for t in p['tasks']});self.assertEqual(p['tasks'][0]['name'],'이름 변경')

    def test_missing_not_deleted(self):
        self.edit(lambda w:w['작업목록'].delete_rows(3))
        plan,p=self.apply();self.assertEqual(plan.missing,['B']);self.assertEqual(len(p['tasks']),2)

    def test_blank_preserve_and_explicit_clear(self):
        self.p['tasks'][0]['note']='보존';self.store.add(self.p,'수정','비고');write_excel(self.file,self.p)
        self.edit(lambda w:setattr(self.cell(w,'note'),'value',None));_,p=self.apply();self.assertEqual(p['tasks'][0]['note'],'보존')
        write_excel(self.file,self.p);self.edit(lambda w:setattr(self.cell(w,'note'),'value',CLEAR));_,p=self.apply();self.assertEqual(p['tasks'][0]['note'],'')

    def test_three_way_conflict_and_local_only_preservation(self):
        self.p['tasks'][0]['name']='앱 변경';self.p['tasks'][1]['note']='앱 전용 변경';self.store.add(self.p,'수정','local')
        self.edit(lambda w:setattr(self.cell(w,'name'),'value','엑셀 변경'))
        plan=prepare(self.file,self.p,self.store);self.assertTrue(any(c.conflict for c in plan.changes))
        with self.assertRaisesRegex(ValueError,'충돌'): materialize(plan,acknowledge=True)
        new=self.store.apply_import(plan,[c.key for c in plan.changes],acknowledge=True)
        self.assertEqual(new['tasks'][0]['name'],'엑셀 변경');self.assertEqual(new['tasks'][1]['note'],'앱 전용 변경')

    def test_duplicate_date_missing_ref_diagnostics(self):
        self.edit(lambda w:setattr(self.cell(w,'id',3),'value','A'))
        plan=prepare(self.file,self.p,self.store);self.assertTrue(any('중복' in str(i) for i in plan.issues))
        with self.assertRaises(ValueError): self.store.apply_import(plan,acknowledge=True)
        write_excel(self.file,self.p)
        self.edit(lambda w:setattr(self.cell(w,'actualStart'),'value','2026-02-30'))
        doc=read_document(self.file);issue=next(i for i in doc.issues if i.severity=='오류')
        self.assertEqual(issue.sheet,'작업목록');self.assertEqual(issue.row,2);self.assertEqual(issue.column,'O');self.assertTrue(issue.fix)
        write_excel(self.file,self.p)
        self.edit(lambda w:setattr(w['관계']['A2'],'value','missing'))
        self.assertTrue(any('존재하지' in str(i) for i in prepare(self.file,self.p,self.store).issues))

    def test_import_atomic_rollback_database_failure(self):
        self.edit(lambda w:setattr(self.cell(w,'name'),'value','새 이름'))
        plan=prepare(self.file,self.p,self.store)
        before=deepcopy(self.store.load(self.p['uid']))
        real=self.store._write
        def fail(p):
            real(p);raise RuntimeError('disk simulation')
        with patch.object(self.store,'_write',side_effect=fail):
            with self.assertRaises(RuntimeError): self.store.apply_import(plan,acknowledge=True)
        self.assertEqual(self.store.load(self.p['uid']),before)
        self.assertEqual(self.store.db.execute('SELECT count(*) FROM versions WHERE project_uid=?',(self.p['uid'],)).fetchone()[0],1)
        self.assertEqual(self.store.db.execute('SELECT status FROM imports ORDER BY id DESC LIMIT 1').fetchone()[0],'failed')
        self.assertTrue(list(self.root.glob('*.pre-import-*')))

    def test_stale_preview_blocks(self):
        plan=prepare(self.file,self.p,self.store)
        self.p['name']='나중 수정';self.store.add(self.p,'수정','새 변경')
        with self.assertRaisesRegex(ValueError,'미리보기'): self.store.apply_import(plan,acknowledge=True)

    def test_selected_subset_revalidated(self):
        def edit(w):
            s=w['작업목록'];values=[c.value for c in s[2]]
            values[list(COLS).index('id')]='C';values[list(COLS).index('uid')]='';s.append(values)
            w['관계'].append(['C','B','FS',0,'달력일','default','',''])
            self.cell(w,'pred',3).value='A, C'
        self.edit(edit);plan=prepare(self.file,self.p,self.store)
        with self.assertRaisesRegex(ValueError,'존재하지'): self.store.apply_import(plan,['$project:relations'],acknowledge=True)

    def test_baseline_forecast_backup_reopen_restore(self):
        r=deterministic(self.p);self.store.baseline(self.p,'승인 1',r);original=self.store.baselines(self.p['uid'])
        backup=self.root/'backup.sqlite3';self.store.backup(backup)
        self.p['tasks'][0].update(o=8,m=8,p=8);self.store.add(self.p,'편집','기간 변경');deterministic(self.p)
        self.assertEqual(self.store.baselines(self.p['uid']),original)
        self.store.db.close();self.store=Store(self.root/'db.sqlite3')
        self.assertEqual(self.store.load(self.p['uid'])['tasks'][0]['m'],8)
        self.store.restore_backup(backup)
        self.assertEqual(self.store.load(self.p['uid'])['tasks'][0]['m'],2)
        self.assertEqual(self.store.baselines(self.p['uid']),original)
        self.assertTrue(self.store.history(self.p['uid']))

    def test_clone_isolation_mapping_and_stable_identity(self):
        p=clone_project(self.p);self.store.add(p,'복제','시험')
        self.assertNotEqual(p['uid'],self.p['uid']);self.assertFalse({t['uid'] for t in p['tasks']} & {t['uid'] for t in self.p['tasks']})
        self.assertEqual(p['relations'][0]['pred'],p['tasks'][0]['id']);self.assertEqual(len(self.store.projects()),2)
        self.assertTrue(all(r[2]==p['id'] for r in self.store.rows(p['uid'])))
        self.assertEqual(self.store.load(self.p['uid'])['tasks'][0]['name'],'작업 A')

    def test_export_formula_safe_and_source_unchanged(self):
        self.p['tasks'][0]['note']='=HYPERLINK("bad")';write_excel(self.file,self.p)
        w=load_workbook(self.file);self.assertEqual(self.cell(w,'note').data_type,'s');w.close()
        self.assertEqual(read_excel(self.file)['tasks'][0]['note'],self.p['tasks'][0]['note'])

    def test_relationship_sheet_precedence(self):
        self.edit(lambda w:setattr(self.cell(w,'pred',3),'value',''))
        plan=prepare(self.file,self.p,self.store)
        self.assertTrue(any('관계 시트' in str(i) and i.severity=='경고' for i in plan.issues))
        self.assertEqual(materialize(plan,acknowledge=True)['relations'],self.p['relations'])

    def test_restore_rejects_non_project_without_damage(self):
        other=self.root/'other.db';sqlite3.connect(other).close();before=self.store.load(self.p['uid'])
        with self.assertRaises(ValueError): self.store.restore_backup(other)
        self.assertEqual(self.store.load(self.p['uid']),before)

    def test_explicit_holiday_clear_and_default_calendar(self):
        self.p['holidays']=['2026-10-05'];self.p['calendars']={'default':dict(weekmask=[0,1,2,3,4],holidays=['2026-10-05'],workdays=[])}
        self.store.add(self.p,'수정','휴일');write_excel(self.file,self.p)
        self.edit(lambda w:setattr(w['휴일']['A2'],'value',CLEAR));_,p=self.apply()
        self.assertEqual(p['holidays'],[]);self.assertEqual(p['calendars']['default']['holidays'],[])

    def test_blank_relation_cells_preserve_values(self):
        self.p['relations'][0].update(lag=2,note='보존 근거');self.store.add(self.p,'수정','관계');write_excel(self.file,self.p)
        self.edit(lambda w:(setattr(w['관계']['D2'],'value',None),setattr(w['관계']['G2'],'value',None)))
        plan,p=self.apply();self.assertEqual(plan.changes,[])
        self.assertEqual(p['relations'][0]['lag'],2);self.assertEqual(p['relations'][0]['note'],'보존 근거')

    def test_idless_new_and_existing_mapping_block(self):
        self.edit(lambda w:(setattr(self.cell(w,'id'),'value',None),setattr(self.cell(w,'uid'),'value',None),w.remove(w['관계']),setattr(self.cell(w,'pred',3),'value',None)))
        plan=prepare(self.file,self.p,self.store)
        self.assertTrue(any('ID 없는' in str(i) for i in plan.issues))
        # Independent external project registration; no name similarity merge.
        w=load_workbook(self.file)
        for row in w['프로젝트'].iter_rows(min_row=2):
            if row[0].value=='프로젝트 UID': row[1].value=None
        w['작업목록'].delete_rows(3);w.save(self.file);w.close()
        plan=prepare(self.file)
        p=materialize(plan,acknowledge=True)
        self.assertEqual(len(p['tasks']),1);self.assertTrue(p['tasks'][0]['id'].startswith('T-'))

    def test_optimistic_write_conflict_no_partial_data(self):
        other=Store(self.root/'db.sqlite3');stale=other.load(self.p['uid'])
        self.p['name']='첫 저장';self.store.add(self.p,'편집','첫 저장')
        stale['name']='늦은 저장'
        try:
            with self.assertRaisesRegex(ValueError,'충돌'): other.add(stale,'편집','충돌')
            self.assertEqual(other.load(self.p['uid'])['name'],'첫 저장')
        finally: other.db.close()


class MigrationTests(unittest.TestCase):
    def test_v1_preserved_and_backup_created(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'legacy.db';db=sqlite3.connect(path)
            db.execute('CREATE TABLE events(seq INTEGER PRIMARY KEY,created TEXT,pj TEXT,kind TEXT,reason TEXT,payload TEXT)')
            p=project();p.pop('uid');p.pop('data_version')
            raw=json.dumps(dict(project=p,summary=None),ensure_ascii=False)
            db.execute('INSERT INTO events VALUES(1,?,?,?,?,?)',('2026-01-01','PJ','기준선','이전 승인 입력',raw));db.commit();db.close()
            store=Store(path)
            self.assertEqual(store.db.execute('SELECT payload FROM events').fetchone()[0],raw)
            self.assertEqual(store.db.execute('PRAGMA user_version').fetchone()[0],2)
            self.assertEqual(len(store.projects()),1);self.assertTrue(store.baselines(store.projects()[0][0])[0]['legacy'])
            self.assertTrue(list(Path(td).glob('*.pre-v2-*')))
            migrated=store.load(store.projects()[0][0]);event=store.get(1)['project']
            self.assertEqual(event['uid'],migrated['uid']);self.assertEqual(event['tasks'][0]['uid'],migrated['tasks'][0]['uid'])
            self.assertTrue(store.history(migrated['uid']));store.db.close()

    def test_failed_migration_rolls_back_schema_and_preserves_original(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'badlegacy.db';db=sqlite3.connect(path)
            db.execute('CREATE TABLE events(seq INTEGER PRIMARY KEY,created TEXT,pj TEXT,kind TEXT,reason TEXT,payload TEXT)')
            db.execute("INSERT INTO events VALUES (1,'t','PJ','x','r','not json')");db.commit();db.close()
            with self.assertRaises(ValueError): Store(path)
            db=sqlite3.connect(path)
            self.assertEqual(db.execute('PRAGMA user_version').fetchone()[0],0)
            self.assertIsNone(db.execute("SELECT name FROM sqlite_master WHERE name='projects'").fetchone())
            self.assertEqual(db.execute('SELECT payload FROM events').fetchone()[0],'not json');db.close()
            self.assertTrue(list(Path(td).glob('*.pre-v2-*')))


if __name__=='__main__': unittest.main(verbosity=2)
