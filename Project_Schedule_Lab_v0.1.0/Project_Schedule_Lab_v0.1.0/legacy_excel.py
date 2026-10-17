"""Python application's user-requested Excel interchange; never overwrites the input file."""
from datetime import date,datetime
from pathlib import Path
import os,tempfile
from openpyxl import load_workbook,Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from engine import summary,percentile,textday
COLS={'id':'Activity ID','module':'모듈','name':'작업명','state':'상태','role':'담당','deliverable':'산출물','calendar':'기간단위','o':'O','m':'M','p':'P','pred':'선행 ID','ro':'잔여 O','rm':'잔여 M','rp':'잔여 P','actualStart':'실제 시작','actualEnd':'실제 완료','constraint':'착수 제한일','baselineStart':'기준 시작','baselineEnd':'기준 완료','note':'비고'}
DATES=['actualStart','actualEnd','constraint','baselineStart','baselineEnd']
def dt(v):
    if isinstance(v,(date,datetime)):return v.strftime('%Y-%m-%d')
    return str(v or '').strip()
def number(v):
    if v in (None,''):return None
    if isinstance(v,str) and v.startswith('='):raise ValueError('기간 입력칸에 수식이 있습니다. 숫자 값으로 붙여넣으세요.')
    return float(v)
def read_excel(path):
    w=load_workbook(path,data_only=False)
    if '작업목록' in w.sheetnames:
        cfg={r[0]:r[1] for r in w['프로젝트'].iter_rows(min_row=2,values_only=True) if r[0]};mods={str(r[0]):str(r[1] or '검토 필요') for r in w['모듈'].iter_rows(min_row=2,values_only=True) if r[0]}
        p={'id':str(cfg.get('PJ_ID') or 'PJ-NEW'),'name':str(cfg.get('프로젝트명') or Path(path).stem),'start':dt(cfg.get('시작일')),'asof':dt(cfg.get('기준일')),'target':dt(cfg.get('목표일')),'holidays':[],'modules':mods,'tasks':[]}
        if '휴일' in w.sheetnames:p['holidays']=[dt(r[0]) for r in w['휴일'].iter_rows(min_row=2,values_only=True) if r[0]]
        rows=list(w['작업목록'].values);headers=list(rows[0]);missing=set(COLS.values())-set(headers)
        if missing:raise ValueError('필수 열 누락: '+', '.join(sorted(missing)))
        for row in rows[1:]:
            if not any(v is not None for v in row):continue
            t={k:row[headers.index(v)] if headers.index(v)<len(row) else None for k,v in COLS.items()}
            for k in ['id','module','name','state','role','deliverable','calendar','note']:t[k]=str(t[k] or '').strip()
            t['pred']=[s.strip() for s in str(t['pred'] or '').replace(';',',').split(',') if s.strip()]
            for k in ['o','m','p','ro','rm','rp']:t[k]=number(t[k])
            for k in DATES:t[k]=dt(t[k])
            t['state']=t['state'] or '계획';t['calendar']=t['calendar'] or '달력일';p['tasks'].append(t)
    elif all(s in w.sheetnames for s in ['상세일정','모듈선택','프로젝트설정']):
        c=w['프로젝트설정'];p={'id':str(c['B6'].value or 'PJ-NEW'),'name':str(c['B9'].value or Path(path).stem),'start':dt(c['B7'].value),'asof':dt(c['B8'].value),'target':dt(c['B11'].value),'holidays':[],'modules':{},'tasks':[]}
        for r in w['모듈선택'].iter_rows(min_row=6,values_only=True):
            if r[0]:p['modules'][str(r[0])]=str(r[1] or '검토 필요')
        for r in w['상세일정'].iter_rows(min_row=6,values_only=True):
            if not r[0]:continue
            t={k:'' for k in COLS};t.update(id=str(r[0]),module=str(r[1]),name=str(r[2]),state=r[4] or '계획',role=r[5] or '',deliverable=r[6] or '',calendar='달력일',pred=[str(x) for x in r[23:29] if x])
            for k,v in zip(['o','m','p','ro','rm','rp'],r[12:18]):t[k]=number(v)
            for k,v in zip(['baselineStart','baselineEnd','actualStart','actualEnd','constraint'],r[7:12]):t[k]=dt(v)
            p['tasks'].append(t)
    else:raise ValueError('지원 형식: 공통 Rev1 또는 이 프로그램에서 내보낸 교환용 XLSX입니다.')
    w.close()
    if len(p['tasks'])>2000:raise ValueError('최대 2,000개 작업을 지원합니다.')
    return p
